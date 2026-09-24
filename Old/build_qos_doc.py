#!/usr/bin/env python3
# build_qos_doc.py
#
# Generates:  LCA-QoS-Rules-2023.docx
# Uses:       python-docx  (pip install python-docx)

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT as ALIGN
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# ----------  Helper functions ----------
def add_heading(text, level):
    """Add a Word heading and return the paragraph object."""
    p = doc.add_heading(text, level=level)
    p.alignment = ALIGN.LEFT
    return p

def add_para(text, style=None):
    """Add a normal paragraph."""
    p = doc.add_paragraph(text, style=style or "Normal")
    p.alignment = ALIGN.JUSTIFY
    return p

def add_table(rows, cols, widths_cm=None, first_row_bold=True):
    """Create a table and optionally set column widths."""
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths_cm:
        for col, width in zip(table.columns, widths_cm):
            for cell in col.cells:
                cell.width = Cm(width)
    if first_row_bold:
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
    return table

# ----------  Title & Gazette Info ----------
add_heading("LESOTHO", level=0)
add_para("Government Gazette\nVol. 68 — Friday 21 April 2023 — No. 29")
doc.add_paragraph()

add_heading("LEGAL NOTICE No. 41 of 2023", level=1)
add_heading("Lesotho Communications Authority (Quality of Service) Rules, 2023", level=2)

# ----------  Arrangement of Rules ----------
add_heading("Arrangement of Rules", level=2)
arrangement = [
    (1, "Citation and commencement"),
    (2, "Scope and application"),
    (3, "Definitions"),
    (4, "Objectives"),
    (5, "Monitoring"),
    (6, "Licensee obligations"),
    (7, "Provision of information"),
    (8, "Investigations, inspections and inquiry"),
    (9, "Sanctions"),
    (10, "Determination of sanctions"),
    (11, "Individual complaints"),
    (12, "Reconsideration"),
    (13, "Repeal"),
]
tbl = add_table(rows=1+len(arrangement), cols=2, widths_cm=[2.0, 12.5])
tbl.style = "Table Grid"
tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Rule", "Title"
for i, (no_, title) in enumerate(arrangement, start=1):
    tbl.cell(i, 0).text = str(no_)
    tbl.cell(i, 1).text = title

doc.add_paragraph()

# ----------  Body Sections ----------
sections = {
    1: "These Rules may be cited as the *Lesotho Communications Authority (Quality of Service) "
       "Rules, 2023* and shall come into operation on the date of publication in the Gazette.",
    2: "These Rules prescribe minimum quality-of-service (QoS) standards applicable to licensees "
       "providing voice, data, postal and courier services.",
    3: "Key definitions:\n"
       "• customer – any person who is or may become an end-user of communications services.\n"
       "• delay – failure to deliver a postal item within the agreed time.\n"
       "• fault – a network state that does not meet service specifications and requires repair.\n"
       "• service accessibility – ability of the network to grant customers the service requested.\n"
       "• service retainability – ability of the network to maintain service once established.\n"
       "(See full list in the Schedules.)",
    4: "Objectives include implementing a QoS framework, improving customer satisfaction, "
       "publishing QoS information and protecting consumer interests.",
    # … add sections 5-13 exactly as needed …
}
for rule_no, text in sections.items():
    add_heading(f"Rule {rule_no}", level=3)
    add_para(text)

# ----------  Example Table from Schedule 1 ----------
add_heading("Schedule 1 — Excerpt of Mobile Voice Parameters", level=2)
voice_table = add_table(rows=7, cols=5,
                        widths_cm=[2.0, 4.0, 5.5, 3.0, 2.0])
headers = ["Ref", "Parameter", "Definition (abridged)", "Target", "Ref-Doc"]
for j, hdr in enumerate(headers):
    voice_table.cell(0, j).text = hdr

data_rows = [
    ("1.1.1", "Call set-up time", "Dial to ring-back", "≤ 8 s", "ETSI EG 202 057-2"),
    ("1.1.2", "Unsuccessful call ratio", "Failed ÷ attempts", "≤ 2 %", "ETSI EG 202 057-3"),
    ("1.1.3", "Call drop ratio", "Dropped ÷ success", "≤ 2 %", "ETSI EG 202 057-3"),
    ("1.1.4", "Network availability", "Core/BTS uptime", "99.99 % / 95 %", "—"),
    ("1.1.5", "Voice quality", "MOS/POLQA", "≥ 3.0", "ITU-T P.863"),
    ("1.1.6", "Coverage (4 G)", "RSRP ≥ -100 dBm", "100 %", "—"),
]
for i, row in enumerate(data_rows, start=1):
    for j, item in enumerate(row):
        voice_table.cell(i, j).text = item

# ----------  Closing ----------
add_para("\nDated 11 April 2023")
add_para("Nizam Goolam\nChief Executive Officer (a.i)\nLesotho Communications Authority")

# ----------  Save ----------
doc.save("LCA-QoS-Rules-2023.docx")
print("Word document created: LCA-QoS-Rules-2023.docx")